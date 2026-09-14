"""
generate_data.py
-----------------
Regenerates data.json for the Qiddiya Car Park dashboard from the three
source Excel files. Run this whenever the source files are updated.

USAGE:
    pip install pandas openpyxl xlrd
    python generate_data.py \
        --submittals "Qiddiya_-_Car_Park_-_Report.xls" \
        --qtycon "Qty_Con.xlsx" \
        --mobilization "Mobilization_WBS_111.xlsx" \
        --out data.json

All three --submittals/--qtycon/--mobilization arguments default to the
filenames above if omitted, so if you keep the same filenames in this folder
you can simply run:
    python generate_data.py
"""
import argparse
import datetime
import json

import pandas as pd
from openpyxl import load_workbook


def build(submittals_path, qtycon_path, mobilization_path):
    # ---- 1. Submittal log, expected targets, S-curve progress ----
    xl = pd.ExcelFile(submittals_path, engine="xlrd" if submittals_path.lower().endswith(".xls") else None)
    db = xl.parse("database Qiddiya")
    exp = xl.parse("Expected number")
    prog = xl.parse("Progress")

    db = db.rename(columns={"Tile": "Zone", "Building Levels": "Level"})
    db["Date Created"] = pd.to_datetime(db["Date Created"])

    disciplines = sorted(db["Discipline"].dropna().unique().tolist())
    doctypes = sorted(db["Document Type"].dropna().unique().tolist())
    statuses = sorted(db["Status"].dropna().unique().tolist())
    zones = sorted(db["Zone"].dropna().unique().tolist())

    d_idx = {v: i for i, v in enumerate(disciplines)}
    t_idx = {v: i for i, v in enumerate(doctypes)}
    s_idx = {v: i for i, v in enumerate(statuses)}
    z_idx = {v: i for i, v in enumerate(zones)}

    epoch = datetime.date(2025, 1, 1)
    rows, docnos, titles = [], [], []
    for _, r in db.iterrows():
        day_off = (r["Date Created"].date() - epoch).days
        rows.append([
            d_idx[r["Discipline"]], t_idx[r["Document Type"]], s_idx[r["Status"]],
            z_idx.get(r["Zone"], -1), day_off,
        ])
        docnos.append(r["Document No"])
        titles.append(str(r["Title"])[:70])

    report_date = db["Date Created"].max().date()
    report_day_off = (report_date - epoch).days

    # Expected/target counts, broken down by discipline (matches the source
    # "Expected number" sheet exactly, so gauge targets stay discipline-aware).
    exp_target_by_discipline = {}
    for _, r in exp.iterrows():
        dt, disc, t = r["Document Type"], r["Discipline"], r["Expected number to be submitted"]
        if pd.isna(t):
            t = 0
        exp_target_by_discipline.setdefault(dt, {})
        exp_target_by_discipline[dt][disc] = exp_target_by_discipline[dt].get(disc, 0) + t

    def clean(v):
        return 0 if pd.isna(v) else float(v)

    progress = {
        "actual": clean(prog["Progress Cum. Actual"][0]),
        "plan": clean(prog["Progress Cum. Plan"][0]),
        "weekly": clean(prog["Weekly Progress.%"][0]),
        "pre": clean(prog["Progress Pre.%"][0]),
    }

    # ---- 2. Concrete quantity survey ----
    qc = pd.read_excel(qtycon_path, sheet_name="Qty Con")
    qc_items = []
    for _, row in qc.iterrows():
        boq = float(row["BOQ Qty"])
        act = float(row["Actual poured Qty"]) if not pd.isna(row["Actual poured Qty"]) else 0.0
        qc_items.append({
            "item": str(row["BOQ Item"]).strip(), "type": row["Con. Type"],
            "boq": boq, "budget": float(row["Budget Qty"]), "actual": act,
        })

    # ---- 3. Site mobilization WBS ----
    wb = load_workbook(mobilization_path, data_only=True)
    ws = wb["0-Main sheet"]
    mob_rows_raw = list(ws.iter_rows(min_row=2, values_only=True))
    mob = pd.DataFrame(
        mob_rows_raw,
        columns=["Area", "Item", "BreakdownStructure", "Delivered", "Installed",
                 "ActualProgressPct", "AssumedWeight", "Achieved", "SubTotal", "Total"],
    )
    mob = mob.dropna(subset=["Area"])

    areas = sorted(mob["Area"].unique().tolist())
    items = sorted(mob["Item"].unique().tolist())
    mob_rows = []
    for _, r in mob.iterrows():
        mob_rows.append({
            "area": r["Area"], "item": r["Item"], "bds": r["BreakdownStructure"],
            "delivered": None if pd.isna(r["Delivered"]) else str(r["Delivered"]),
            "installed": None if pd.isna(r["Installed"]) else str(r["Installed"]),
            "weight": float(r["AssumedWeight"]) if not pd.isna(r["AssumedWeight"]) else 0.0,
            "achieved": float(r["Achieved"]) if not pd.isna(r["Achieved"]) else 0.0,
        })

    data = {
        "epoch": epoch.isoformat(),
        "reportDayOff": report_day_off,
        "reportDateStr": report_date.strftime("%d %B %Y"),
        "disciplines": disciplines, "doctypes": doctypes, "statuses": statuses, "zones": zones,
        "rows": rows, "docnos": docnos, "titles": titles,
        "expTargetByDiscipline": exp_target_by_discipline,
        "progress": progress,
        # Concrete cutoff date is informational only — update it here if the
        # QS cutoff date changes between reports.
        "concreteCutoff": "27 July 2026",
        "qcItems": qc_items,
        "areas": areas, "items": items, "mobRows": mob_rows,
    }
    return data


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--submittals", default="Qiddiya_-_Car_Park_-_Report.xls")
    parser.add_argument("--qtycon", default="Qty_Con.xlsx")
    parser.add_argument("--mobilization", default="Mobilization_WBS_111.xlsx")
    parser.add_argument("--out", default="data.json")
    args = parser.parse_args()

    result = build(args.submittals, args.qtycon, args.mobilization)
    with open(args.out, "w", encoding="utf-8") as f:
        json.dump(result, f, allow_nan=False)
    print(f"Wrote {args.out} ({len(result['rows'])} submittal rows, "
          f"{len(result['qcItems'])} concrete items, {len(result['mobRows'])} mobilization rows)")
